#!/usr/bin/env python3
"""أوامر Excel وWord وPowerPoint وPDF للمكتب، بدون ما يفتح أي برنامج.

ما يصرف توكنز وما يحتاج شاشة. يستدعى من agent.ps1 (أمر pc office):
يقرا طلب JSON من ملف ويطبع نتيجة JSON بسطر واحد.
  python3 office.py <request.json>

الطلب: {"op": "...", "path": "...", ...}. المسارات يتأكد منها agent.ps1 قبل ما توصل هنا،
والكتابة تنحفظ نسختها القديمة بـ _نسخ_قبل_التعديل. ما يمسح أي ملف.

Excel:
  xlsx_sheets  path                                   أسماء الأوراق وعدد صفوفها
  xlsx_read    path [sheet] [max] [data_only]         صفوف الورقة
  xlsx_find    path text [sheet]                      وين موجود نص (خلية وورقة)
  xlsx_set     path cells {"A1": "..", "B2": "=..."}  يكتب خلايا ومعادلات
  xlsx_append  path row [..] [sheet]                  يضيف صف بالآخر
  xlsx_new     path [sheet] [rows [[..], ..]]          ملف جديد
  xlsx_add_sheet path sheet                           ورقة جديدة
Word:
  docx_read    path [max]                             نص المستند
  docx_tables  path                                   جداول المستند
  docx_replace path repl {"قديم": "جديد"}             استبدال نص (يحافظ على التنسيق)
  docx_new     path paras ["..", ..]                  مستند جديد
  docx_append  path paras ["..", ..] | rows [[..]]    يضيف فقرات أو جدول بالآخر
PowerPoint:
  pptx_read    path                                   نص كل سلايد
  pptx_replace path repl {"قديم": "جديد"}             استبدال نص بكل السلايدات
  pptx_new     path slides [{"title": "..", "body": ["..", ..]}, ..]
  pptx_add     path slides [..]                        يضيف سلايدات بالآخر
PDF:
  pdf_text     path [max] [pages "1-3"]               نص الـ PDF (النص المحفور بس، مو المصوّر)
  pdf_info     path                                   عدد الصفحات والقياسات
  pdf_merge    paths [".."] out                       يجمع ملفات بملف واحد
  pdf_split    path out_dir [pages "1-3"]             يفصل الصفحات ملفات
  pdf_rotate   path out angle [pages "1-3"]           يدوّر صفحات
"""
import json
import os
import sys


# ------------------------------------------------------------------ مساعدات

def _ws(wb, sheet):
    if sheet:
        if sheet not in wb.sheetnames:
            raise ValueError("ماكو ورقة باسم: " + sheet)
        return wb[sheet]
    return wb.active


def _pages(spec, total):
    """"1-3" أو "2" أو "1,4-5" → أرقام صفحات (تبدي من 1). فارغ = كل الصفحات."""
    if not spec:
        return list(range(1, total + 1))
    out = []
    for part in str(spec).split(","):
        part = part.strip()
        if not part:
            continue
        if "-" in part:
            a, b = part.split("-", 1)
            out.extend(range(int(a), int(b) + 1))
        else:
            out.append(int(part))
    bad = [p for p in out if p < 1 or p > total]
    if bad:
        raise ValueError("رقم صفحة غلط: %s (الملف %d صفحة)" % (bad[0], total))
    return out


def _limit(obj, name, mx):
    if not isinstance(obj, dict) or len(obj) > mx:
        raise ValueError("%s لازم يكون كائن فيه %d أو أقل" % (name, mx))


# ------------------------------------------------------------------ Excel

def xlsx_sheets(req):
    import openpyxl
    wb = openpyxl.load_workbook(req["path"], read_only=True)
    return {"sheets": [{"name": ws.title, "rows": ws.max_row, "cols": ws.max_column} for ws in wb.worksheets]}


def xlsx_read(req):
    import openpyxl
    wb = openpyxl.load_workbook(req["path"], data_only=bool(req.get("data_only", True)))
    ws = _ws(wb, req.get("sheet"))
    mx = int(req.get("max", 200))
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(["" if v is None else v for v in r])
        if len(rows) >= mx:
            break
    return {"sheet": ws.title, "rows": rows, "dims": ws.dimensions}


def xlsx_find(req):
    import openpyxl
    wb = openpyxl.load_workbook(req["path"], data_only=True)
    needle = str(req["text"])
    sheets = [_ws(wb, req["sheet"])] if req.get("sheet") else wb.worksheets
    hits = []
    for ws in sheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None and needle in str(cell.value):
                    hits.append({"sheet": ws.title, "cell": cell.coordinate, "value": str(cell.value)[:200]})
                    if len(hits) >= 100:
                        return {"hits": hits, "more": True}
    return {"hits": hits}


def xlsx_set(req):
    import openpyxl
    wb = openpyxl.load_workbook(req["path"])
    ws = _ws(wb, req.get("sheet"))
    cells = req["cells"]
    _limit(cells, "cells", 500)
    for ref, val in cells.items():
        ws[ref] = val
    wb.save(req["path"])
    return {"sheet": ws.title, "written": len(cells)}


def xlsx_append(req):
    import openpyxl
    wb = openpyxl.load_workbook(req["path"])
    ws = _ws(wb, req.get("sheet"))
    rows = req.get("rows")
    if rows is None:
        rows = [req["row"]]
    if not isinstance(rows, list) or len(rows) > 500:
        raise ValueError("rows لازم قائمة 500 صف أو أقل")
    for r in rows:
        ws.append(r)
    wb.save(req["path"])
    return {"sheet": ws.title, "added": len(rows), "last_row": ws.max_row}


def xlsx_new(req):
    import openpyxl
    if os.path.exists(req["path"]):
        raise ValueError("الملف موجود أصلاً: " + req["path"])
    wb = openpyxl.Workbook()
    ws = wb.active
    if req.get("sheet"):
        ws.title = req["sheet"]
    for r in req.get("rows", []):
        ws.append(r)
    wb.save(req["path"])
    return {"path": req["path"], "sheet": ws.title, "rows": ws.max_row}


def xlsx_add_sheet(req):
    import openpyxl
    wb = openpyxl.load_workbook(req["path"])
    name = req["sheet"]
    if name in wb.sheetnames:
        raise ValueError("الورقة موجودة أصلاً: " + name)
    ws = wb.create_sheet(name)
    for r in req.get("rows", []):
        ws.append(r)
    wb.save(req["path"])
    return {"sheet": name, "sheets": wb.sheetnames}


# ------------------------------------------------------------------ Word

def docx_read(req):
    import docx
    doc = docx.Document(req["path"])
    mx = int(req.get("max", 400))
    return {"paragraphs": [p.text for p in doc.paragraphs if p.text.strip()][:mx]}


def docx_tables(req):
    import docx
    doc = docx.Document(req["path"])
    tables = []
    for t in doc.tables[:20]:
        tables.append([[c.text for c in row.cells] for row in t.rows[:100]])
    return {"tables": tables}


def docx_replace(req):
    import docx
    doc = docx.Document(req["path"])
    repl = req["repl"]
    _limit(repl, "repl", 200)
    n = 0

    def fix(paras):
        nonlocal n
        for p in paras:
            for run in p.runs:
                for old, new in repl.items():
                    if old in run.text:
                        run.text = run.text.replace(old, str(new))
                        n += 1

    fix(doc.paragraphs)
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                fix(cell.paragraphs)
    doc.save(req["path"])
    return {"replaced": n}


def docx_new(req):
    import docx
    if os.path.exists(req["path"]):
        raise ValueError("الملف موجود أصلاً: " + req["path"])
    doc = docx.Document()
    for text in req.get("paras", []):
        doc.add_paragraph(str(text))
    doc.save(req["path"])
    return {"path": req["path"], "paragraphs": len(req.get("paras", []))}


def docx_append(req):
    import docx
    doc = docx.Document(req["path"])
    added = 0
    for text in req.get("paras", []):
        doc.add_paragraph(str(text))
        added += 1
    rows = req.get("rows")
    if rows:
        if not isinstance(rows, list) or not rows:
            raise ValueError("rows لازم قائمة صفوف")
        t = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, v in enumerate(r):
                t.cell(i, j).text = "" if v is None else str(v)
        added += len(rows)
    doc.save(req["path"])
    return {"added": added}


# ------------------------------------------------------------------ PowerPoint

def pptx_read(req):
    from pptx import Presentation
    prs = Presentation(req["path"])
    slides = []
    for i, s in enumerate(prs.slides, 1):
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        slides.append({"slide": i, "text": texts})
    return {"slides": slides, "count": len(slides)}


def pptx_replace(req):
    from pptx import Presentation
    prs = Presentation(req["path"])
    repl = req["repl"]
    _limit(repl, "repl", 200)
    n = 0
    for s in prs.slides:
        for sh in s.shapes:
            if not sh.has_text_frame:
                continue
            for p in sh.text_frame.paragraphs:
                for run in p.runs:
                    for old, new in repl.items():
                        if old in run.text:
                            run.text = run.text.replace(old, str(new))
                            n += 1
    prs.save(req["path"])
    return {"replaced": n}


def _add_slides(prs, slides):
    layout = prs.slide_layouts[1]  # عنوان ومحتوى
    for item in slides:
        slide = prs.slides.add_slide(layout)
        slide.shapes.title.text = str(item.get("title", ""))
        body = item.get("body") or []
        if body:
            tf = slide.placeholders[1].text_frame
            tf.text = str(body[0])
            for line in body[1:]:
                tf.add_paragraph().text = str(line)
    return len(slides)


def pptx_new(req):
    from pptx import Presentation
    if os.path.exists(req["path"]):
        raise ValueError("الملف موجود أصلاً: " + req["path"])
    prs = Presentation()
    slides = req.get("slides", [])
    if len(slides) > 100:
        raise ValueError("100 سلايد أو أقل")
    n = _add_slides(prs, slides)
    prs.save(req["path"])
    return {"path": req["path"], "slides": n}


def pptx_add(req):
    from pptx import Presentation
    prs = Presentation(req["path"])
    slides = req.get("slides", [])
    if len(slides) > 100:
        raise ValueError("100 سلايد أو أقل")
    n = _add_slides(prs, slides)
    prs.save(req["path"])
    return {"added": n, "total": len(prs.slides._sldIdLst)}


# ------------------------------------------------------------------ PDF

def pdf_info(req):
    from pypdf import PdfReader
    r = PdfReader(req["path"])
    return {"pages": len(r.pages), "encrypted": r.is_encrypted,
            "size": [[round(float(p.mediabox.width)), round(float(p.mediabox.height))] for p in r.pages[:20]]}


def pdf_text(req):
    from pypdf import PdfReader
    r = PdfReader(req["path"])
    want = _pages(req.get("pages"), len(r.pages))
    mx = int(req.get("max", 20000))
    out, total = [], 0
    for p in want:
        t = (r.pages[p - 1].extract_text() or "").strip()
        total += len(t)
        out.append({"page": p, "text": t[:4000]})
        if total >= mx:
            break
    empty = all(not o["text"] for o in out)
    return {"pages": out, "scanned": empty,
            "note": "الـ PDF مصوّر (ماكو نص محفور)، يحتاج قراية كصورة" if empty else ""}


def pdf_merge(req):
    from pypdf import PdfWriter
    paths = req["paths"]
    if not isinstance(paths, list) or len(paths) < 2 or len(paths) > 100:
        raise ValueError("paths لازم من 2 لـ 100 ملف")
    w = PdfWriter()
    for p in paths:
        w.append(p)
    with open(req["out"], "wb") as fh:
        w.write(fh)
    return {"out": req["out"], "files": len(paths), "pages": len(w.pages)}


def pdf_split(req):
    from pypdf import PdfReader, PdfWriter
    r = PdfReader(req["path"])
    want = _pages(req.get("pages"), len(r.pages))
    out_dir = req["out_dir"]
    os.makedirs(out_dir, exist_ok=True)
    base = os.path.splitext(os.path.basename(req["path"]))[0]
    made = []
    for p in want:
        w = PdfWriter()
        w.add_page(r.pages[p - 1])
        dst = os.path.join(out_dir, "%s - صفحة %d.pdf" % (base, p))
        with open(dst, "wb") as fh:
            w.write(fh)
        made.append(dst)
    return {"files": made, "count": len(made)}


def pdf_rotate(req):
    from pypdf import PdfReader, PdfWriter
    angle = int(req.get("angle", 90))
    if angle % 90:
        raise ValueError("الزاوية لازم 90 أو 180 أو 270")
    r = PdfReader(req["path"])
    want = set(_pages(req.get("pages"), len(r.pages)))
    w = PdfWriter()
    for i, page in enumerate(r.pages, 1):
        if i in want:
            page.rotate(angle)
        w.add_page(page)
    with open(req["out"], "wb") as fh:
        w.write(fh)
    return {"out": req["out"], "rotated": len(want), "angle": angle}


OPS = {
    "xlsx_sheets": xlsx_sheets, "xlsx_read": xlsx_read, "xlsx_find": xlsx_find,
    "xlsx_set": xlsx_set, "xlsx_append": xlsx_append, "xlsx_new": xlsx_new,
    "xlsx_add_sheet": xlsx_add_sheet,
    "docx_read": docx_read, "docx_tables": docx_tables, "docx_replace": docx_replace,
    "docx_new": docx_new, "docx_append": docx_append,
    "pptx_read": pptx_read, "pptx_replace": pptx_replace, "pptx_new": pptx_new, "pptx_add": pptx_add,
    "pdf_info": pdf_info, "pdf_text": pdf_text, "pdf_merge": pdf_merge,
    "pdf_split": pdf_split, "pdf_rotate": pdf_rotate,
}
# الأوامر اللي تكتب: agent.ps1 يحفظ نسخة قديمة قبلها
WRITES = ["xlsx_set", "xlsx_append", "xlsx_add_sheet", "docx_replace", "docx_append",
          "pptx_replace", "pptx_add"]


def main(argv):
    if not argv:
        sys.exit("يحتاج ملف الطلب")
    with open(argv[0], encoding="utf-8") as fh:
        req = json.load(fh)
    op = req.get("op")
    if op not in OPS:
        print(json.dumps({"ok": False, "error": "أمر office غير معروف: " + str(op)}, ensure_ascii=False))
        return
    try:
        data = OPS[op](req)
        print(json.dumps(dict(data, ok=True), ensure_ascii=False, default=str))
    except Exception as e:  # noqa: BLE001
        print(json.dumps({"ok": False, "error": "%s: %s" % (type(e).__name__, e)}, ensure_ascii=False))


if __name__ == "__main__":
    main(sys.argv[1:])
