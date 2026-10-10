"""أدوات Word: نبني كل كتاب فوق كتاب حقيقي سابق للشركة (نموذج الكتاب).

الفكرة: ما نصمم ترويسة ولا نختار خطوط من عندنا. ناخذ كتاب معتمد للشركة
(بيه الترويسة والختم والخطوط والمحاذاة اللي يحبها حارث)، نلكه بيه:
    سطر «الى /»، سطر «م/»، فقرة المتن، سطر الختام
ونستنسخ تنسيقها للفقرات الجديدة. الترويسة والتوقيع والإيميل يبقون مثل ما هم.
"""
import copy
import re

from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TO_PREFIXES = ('الى', 'إلى', 'الي')
SUBJECT_RE = re.compile(r'^م\s*/')
CLOSING_PREFIXES = ('وتقبلوا', 'ولكم', 'مع فائق', 'ودمتم', 'وتفضلوا', 'مع التقدير', 'مع الشكر')
NUMBER_RE = re.compile(r'(العدد\s*[:：/]?\s*)([0-9٠-٩\-–/ ]*)')


class TemplateError(Exception):
    pass


def ptext(p):
    return ''.join(t.text or '' for t in p.iter(qn('w:t'))).strip()


def _first_rpr(p):
    for r in p.iter(qn('w:r')):
        if any((t.text or '').strip() for t in r.iter(qn('w:t'))):
            rpr = r.find(qn('w:rPr'))
            return copy.deepcopy(rpr) if rpr is not None else None
    return None


def make_para(proto, text, bold=None, size=None):
    """فقرة جديدة بنفس تنسيق proto ونص واحد."""
    p = copy.deepcopy(proto)
    rpr = _first_rpr(proto)
    for child in list(p):
        if child.tag != qn('w:pPr'):
            p.remove(child)
    r = OxmlElement('w:r')
    if rpr is None:
        rpr = OxmlElement('w:rPr')
    _ensure_rtl(rpr)
    if bold is not None:
        for tag in ('w:b', 'w:bCs'):
            el = rpr.find(qn(tag))
            if bold and el is None:
                rpr.append(OxmlElement(tag))
            elif not bold and el is not None:
                rpr.remove(el)
    if size:
        for tag in ('w:sz', 'w:szCs'):
            el = rpr.find(qn(tag))
            if el is None:
                el = OxmlElement(tag)
                rpr.append(el)
            el.set(qn('w:val'), str(int(size * 2)))
    r.append(rpr)
    t = OxmlElement('w:t')
    t.set(qn('xml:space'), 'preserve')
    t.text = text
    r.append(t)
    p.append(r)
    return p


def _ensure_rtl(rpr):
    if rpr.find(qn('w:rtl')) is None:
        rpr.append(OxmlElement('w:rtl'))


def _set_center(p):
    ppr = p.find(qn('w:pPr'))
    if ppr is None:
        ppr = OxmlElement('w:pPr')
        p.insert(0, ppr)
    jc = ppr.find(qn('w:jc'))
    if jc is None:
        jc = OxmlElement('w:jc')
        ppr.append(jc)
    jc.set(qn('w:val'), 'center')
    return p


class LetterTemplate:
    """كتاب سابق للشركة نستعمله قالب."""

    def __init__(self, path):
        self.path = path
        doc = Document(path)
        body = doc.element.body
        kids = [k for k in body if k.tag in (qn('w:p'), qn('w:tbl'))]
        texts = [ptext(k) if k.tag == qn('w:p') else None for k in kids]

        def find(pred, start=0):
            for i in range(start, len(kids)):
                if texts[i] is not None and pred(texts[i]):
                    return i
            return -1

        i_to = find(lambda t: t.startswith(TO_PREFIXES))
        if i_to < 0:
            raise TemplateError('نموذج الكتاب ما بيه سطر «الى /»: ' + path)
        i_subj = find(lambda t: bool(SUBJECT_RE.match(t)), i_to + 1)
        i_close = find(lambda t: t.startswith(CLOSING_PREFIXES), max(i_to, i_subj) + 1)
        if i_close < 0:
            i_close = find(lambda t: 'المدير' in t, i_to + 1)
            if i_close < 0:
                raise TemplateError('نموذج الكتاب ما بيه سطر ختام ولا «المدير المفوض»: ' + path)
            self.has_closing = False
        else:
            self.has_closing = True
        start_body = (i_subj if i_subj >= 0 else i_to) + 1
        bodies = [i for i in range(start_body, i_close) if texts[i] and len(texts[i]) > 40]
        self.proto = {
            'to': kids[i_to],
            'subject': kids[i_subj] if i_subj >= 0 else _set_center(copy.deepcopy(kids[i_to])),
            'body': kids[bodies[0]] if bodies else kids[i_to],
            'closing': kids[i_close] if self.has_closing else _set_center(copy.deepcopy(kids[i_to])),
            'blank': next((kids[i] for i in range(i_to + 1, i_close) if texts[i] == ''), None),
        }
        self.proto = {k: copy.deepcopy(v) if v is not None else None for k, v in self.proto.items()}
        self.i_to, self.i_close = i_to, i_close

    def new_document(self):
        """نسخة جديدة من النموذج ويه موضع الإدراج (العناصر بين «الى» والختام تنشال)."""
        doc = Document(self.path)
        body = doc.element.body
        kids = [k for k in body if k.tag in (qn('w:p'), qn('w:tbl'))]
        anchor = kids[self.i_close]
        drop = kids[self.i_to:self.i_close + (1 if self.has_closing else 0)]
        if not self.has_closing:
            anchor = kids[self.i_close]
        else:
            anchor = kids[self.i_close + 1] if self.i_close + 1 < len(kids) else None
        for k in drop:
            body.remove(k)
        return doc, anchor

    def strip_header_block(self, doc):
        """للأوراق اللي مو كتب (غلاف، فهرس، ملصقات): نشيل سطور العدد والتاريخ قبل «الى»."""
        body = doc.element.body
        for k in list(body):
            if k.tag != qn('w:p'):
                continue
            t = ptext(k)
            if t.startswith('العدد') or t.startswith('التاريخ'):
                body.remove(k)


def insert_before(doc, anchor, el):
    body = doc.element.body
    if anchor is None:
        sect = body.find(qn('w:sectPr'))
        if sect is not None:
            sect.addprevious(el)
        else:
            body.append(el)
    else:
        anchor.addprevious(el)


def set_number(doc, number):
    """يبدل رقم «العدد» بالمتن والترويسة (حتى لو الرقم مقسوم على أكثر من run)."""
    parts = [doc.element.body]
    for s in doc.sections:
        for hf in (s.header, s.first_page_header, s.footer, s.first_page_footer):
            try:
                parts.append(hf._element)
            except Exception:
                pass
    done = False
    for root in parts:
        for p in root.iter(qn('w:p')):
            ts = list(p.iter(qn('w:t')))
            full = ''.join(t.text or '' for t in ts)
            if 'العدد' not in full:
                continue
            new = NUMBER_RE.sub(lambda m: m.group(1) + str(number) + '  ', full, count=1)
            if new != full and ts:
                ts[0].text = new
                ts[0].set(qn('xml:space'), 'preserve')
                for t in ts[1:]:
                    t.text = ''
                done = True
    return done


def doc_text(doc):
    out = [ptext(p) for p in doc.element.body.iter(qn('w:p'))]
    return '\n'.join(out)


# ------------------------------------------------------------------ الجداول

def _border_xml(tbl):
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement('w:' + side)
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '6')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '444444')
        borders.append(b)
    tblPr.append(borders)
    bidi = OxmlElement('w:bidiVisual')
    tblPr.append(bidi)
    w = OxmlElement('w:tblW')
    w.set(qn('w:w'), '5000')
    w.set(qn('w:type'), 'pct')
    tblPr.append(w)
    jc = OxmlElement('w:jc')
    jc.set(qn('w:val'), 'center')
    tblPr.append(jc)


def _fixed_grid(tbl, widths):
    tblPr = tbl.find(qn('w:tblPr'))
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)
    grid = tbl.find(qn('w:tblGrid'))
    if grid is not None:
        for gc, w in zip(grid.findall(qn('w:gridCol')), widths):
            gc.set(qn('w:w'), str(w))


def _cell_width(tc, w):
    tcPr = tc.get_or_add_tcPr()
    el = tcPr.find(qn('w:tcW'))
    if el is None:
        el = OxmlElement('w:tcW')
        tcPr.insert(0, el)
    el.set(qn('w:w'), str(w))
    el.set(qn('w:type'), 'dxa')


def _shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


def make_table(doc, proto, rows, header_fill='D9E2F3', bold_last=False, size=13, widths=None):
    """جدول من صفوف نصوص. الصف الأول عناوين. proto = فقرة نستنسخ خطها."""
    t = doc.add_table(rows=len(rows), cols=len(rows[0]))
    tbl = t._tbl
    _border_xml(tbl)
    if widths is None:
        weight = {'ت': 0.45, 'المادة': 3.2, 'المتطلب': 5, 'المحتوى': 5, 'الحالة': 2.5, 'الوحدة': 1.1,
                  'الكمية': 1.5, 'الصفحة': 1}
        ws = [weight.get(str(h), 2.8 if 'المبلغ' in str(h) else 2) for h in rows[0]]
        sec = doc.sections[0]
        try:  # العرض المتاح بين الهوامش (twips)
            total_w = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)
        except Exception:
            total_w = 9400
        widths = [int(total_w * w / sum(ws)) for w in ws]
    _fixed_grid(tbl, widths)
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            cell = t.cell(ri, ci)
            tc = cell._tc
            for p in list(tc.iter(qn('w:p'))):
                tc.remove(p)
            bold = ri == 0 or (bold_last and ri == len(rows) - 1)
            para = _set_center(make_para(proto, str(val), bold=bold, size=size))
            _no_indent(para)
            tc.append(para)
            if ri == 0:
                _shade(cell, header_fill)
            elif bold_last and ri == len(rows) - 1:
                _shade(cell, 'F2F2F2')
            _cell_width(tc, widths[ci])
    # add_table يحط الجدول بآخر المتن؛ المتصل ينقله لمكانه
    tbl.getparent().remove(tbl)
    return tbl


def _no_indent(p):
    ppr = p.find(qn('w:pPr'))
    if ppr is None:
        return
    for tag in ('w:ind', 'w:numPr'):
        el = ppr.find(qn(tag))
        if el is not None:
            ppr.remove(el)
    sp = ppr.find(qn('w:spacing'))
    if sp is None:
        sp = OxmlElement('w:spacing')
        ppr.append(sp)
    sp.set(qn('w:before'), '40')
    sp.set(qn('w:after'), '40')


def page_break(proto_blank):
    p = OxmlElement('w:p')
    r = OxmlElement('w:r')
    br = OxmlElement('w:br')
    br.set(qn('w:type'), 'page')
    r.append(br)
    p.append(r)
    return p
